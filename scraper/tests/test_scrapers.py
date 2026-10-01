"""End-to-end tests of the extraction machinery against local fixture sites.

The fixtures mimic common layouts (card grid, JS + JSON API with pagination,
paginated HTML table, cascading dropdowns). They use fictional names and do
not reflect the real government sites' markup.
"""

from __future__ import annotations

import http.server
import json
import sys
import threading
from functools import partial
from pathlib import Path
from urllib.parse import unquote

import pytest

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))

import barangay  # noqa: E402
import common  # noqa: E402
import house  # noqa: E402
import lgu  # noqa: E402
import senate  # noqa: E402
from common import Browser, extract_cards, looks_like_name, strip_honorific  # noqa: E402
from officials import classify, rows_to_officials  # noqa: E402

FIXTURES = HERE / "fixtures"


class Handler(http.server.SimpleHTTPRequestHandler):
    def do_GET(self):  # noqa: N802
        if self.path.startswith("/hq/senators/congress/28"):
            self.path = "/hq/senators/congress/28.json"
        if self.path.startswith("/api/members"):
            page = self.path.split("page=")[-1]
            self.path = f"/api/members_page{page}.json"
        if self.path.startswith("/img/"):
            self.send_response(404)
            self.end_headers()
            return
        super().do_GET()

    def end_headers(self):
        if self.path.endswith(".json"):
            self.send_header("Content-Type", "application/json")
        super().end_headers()

    def log_message(self, *args):
        pass


@pytest.fixture(scope="module")
def site():
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), partial(Handler, directory=str(FIXTURES)))
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()


@pytest.fixture(scope="module")
def browser():
    b = Browser(delay=0.1)
    yield b
    b.close()


@pytest.fixture(autouse=True)
def isolated_output(tmp_path, monkeypatch):
    monkeypatch.setattr(common, "SNAPSHOT_DIR", tmp_path / "snap")
    monkeypatch.setattr(common, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(barangay, "OUT_DIR", tmp_path / "data" / "barangay")
    monkeypatch.setattr(barangay, "STATE", tmp_path / "state.json")
    return tmp_path


def test_name_helpers():
    assert looks_like_name("Juan A. Dela Cruz")
    assert looks_like_name("Ma. Teresa B. Santos-Reyes")
    assert not looks_like_name("Home")
    assert not looks_like_name("About the Senate of the Philippines and its many committees")
    assert not looks_like_name("Tel: 8123-4567")
    assert strip_honorific("Hon. Sen. Juan Dela Cruz") == "Juan Dela Cruz"


def test_classify_positions():
    assert classify("Punong Barangay")[0] == "punong_barangay"
    assert classify("Vice Mayor")[0] == "vice_mayor"
    assert classify("SK Chairperson")[0] == "sk_chair"
    assert classify("Barangay Kagawad")[0] == "kagawad"
    assert classify("City Mayor")[0] == "mayor"
    assert classify("Sangguniang Kabataan Chairperson")[0] == "sk_chair"
    assert classify("SK Kagawad")[0] == "sk_kagawad"
    assert classify("Sangguniang Barangay Member")[0] == "kagawad"
    assert classify("Sangguniang Panlungsod Member")[0] == "councilor"


def test_wide_and_long_rows():
    wide = [{"lgu": "Pateros", "mayor": "Hon. Juan Dela Cruz", "vice_mayor": "Maria Santos", "address": "Pateros"}]
    recs = rows_to_officials(wide, source="s", source_url="u", defaults={"region": "NCR"})
    assert {(r["name"], r["level"], r["lgu"]) for r in recs} == {
        ("Juan Dela Cruz", "mayor", "Pateros"), ("Maria Santos", "vice_mayor", "Pateros")}
    long = [{"position": "Punong Barangay", "name": "Pedro Reyes"}]
    recs = rows_to_officials(long, source="s", source_url="u", defaults={"barangay": "Cogon", "lgu": "X"})
    assert recs[0]["level"] == "punong_barangay" and recs[0]["barangay"] == "Cogon"


def test_senate_api_with_menu_noise(site, browser, monkeypatch):
    """Mirrors the live page: a 31-link menu, heading cards, and a JSON API."""
    monkeypatch.setattr(senate, "URL", f"{site}/senate.html")
    recs = senate.scrape(browser)
    assert len(recs) == 24  # inactive senator excluded, menu links ignored
    by_name = {r["name"]: r for r in recs}
    first = by_name["Juan Dela Cruz"]
    assert first["position"] == "Senator · Senate President"
    assert first["photo"].endswith("/hq/uploads/500.jpg")
    assert first["profile_url"].endswith("/senator/Juan-Dela-Cruz")
    assert "Fictional test record number 0" in first["details"]["biography"]
    quoted = by_name['Maria "Ma" B. Santos']
    assert unquote(quoted["profile_url"]).endswith('/senator/Maria-"Ma"-B.-Santos')
    assert quoted["photo"]


def test_senate_cards_fallback(site, browser, monkeypatch):
    monkeypatch.setattr(senate, "URL", f"{site}/senate_cards.html")
    recs = senate.scrape(browser)
    assert len(recs) == 24
    first = recs[0]
    assert first["name"].count(" ") >= 2 and not first["name"].startswith("Sen.")
    assert first["position"] == "Senator · Senate President"
    assert first["profile_url"].endswith("/senators/sen0")
    assert first["photo"].endswith("/img/sen0.jpg")
    assert first["party"] == "NP"


def test_house_json_api_with_pagination(site, browser, monkeypatch):
    monkeypatch.setattr(house, "URL", f"{site}/house.html")
    recs = house.scrape(browser)
    assert len(recs) == 300
    district = [r for r in recs if r["position"] == "District Representative"]
    partylist = [r for r in recs if r["position"] == "Party-list Representative"]
    assert len(district) == 250 and len(partylist) == 50
    assert district[0]["province"] == "Cebu"
    assert district[0]["party"] == "NUP"


def test_lgu_table_pagination(site, browser, monkeypatch):
    monkeypatch.setattr(lgu, "URL", f"{site}/lgu.html")
    lgus, officials = lgu.scrape(browser)
    assert len(lgus) == 17
    assert len(officials) == 34
    pateros = [o for o in officials if o["lgu"] == "Pateros"]
    assert {o["level"] for o in pateros} == {"mayor", "vice_mayor"}
    assert all(o["region"] == "NCR" for o in officials)


def test_barangay_cascading_selects_and_resume(site, browser, monkeypatch, isolated_output):
    monkeypatch.setattr(barangay, "URL", f"{site}/barangay.html")
    import time
    n = barangay.scrape(browser, deadline=time.time() + 120, only_region=None, reset=True)
    assert n == 5 * 4  # 5 barangays x 4 officials
    index = json.loads((isolated_output / "data" / "barangay" / "index.json").read_text())
    assert index["count"] == 20
    assert {f["province"] for f in index["files"]} == {"Bohol", "Benguet"}
    bohol = json.loads((isolated_output / "data" / "barangay" / "region-vii--bohol.json").read_text())
    cogon = [r for r in bohol["records"] if r["barangay"] == "Cogon"]
    assert {r["level"] for r in cogon} == {"punong_barangay", "kagawad", "sk_chair"}
    assert all(r["lgu"] == "Tagbilaran City" for r in cogon)

    # A second run resumes and finds nothing left to do.
    assert barangay.scrape(browser, deadline=time.time() + 120, only_region=None, reset=False) == 0


def test_write_dataset_rejects_bad_counts(isolated_output):
    with pytest.raises(SystemExit):
        common.write_dataset("senate", [], "u", min_count=20)
    assert not (isolated_output / "data" / "senate.json").exists()


def test_waits_past_bot_check_page(site, browser, monkeypatch):
    """congress.gov.ph shows a Cloudflare check page before the real one."""
    monkeypatch.setattr(house, "URL", f"{site}/challenge.html")
    recs = house.scrape(browser)
    assert len(recs) == 300


# --------------------------------------------------------------------------- #
# Importing hand-saved copies (data-inbox/)
# --------------------------------------------------------------------------- #

def test_import_saved_copies(site, browser, monkeypatch, isolated_output):
    import csv as _csv

    import import_saved
    import run_all
    from openpyxl import Workbook

    inbox = isolated_output / "data-inbox"
    monkeypatch.setattr(import_saved, "INBOX", inbox)
    monkeypatch.setattr(run_all, "DATA_DIR", isolated_output / "data")

    # House: a rendered page saved by the browser, as MHTML ("Download page")
    # for page 1 and plain HTML ("Save page as") for page 2.
    (inbox / "house").mkdir(parents=True)
    browser.page.goto(f"{site}/house.html")
    browser.page.wait_for_selector(".member")
    cdp = browser.page.context.new_cdp_session(browser.page)
    mhtml = cdp.send("Page.captureSnapshot", {"format": "mhtml"})["data"]
    (inbox / "house" / "page1.mhtml").write_text(mhtml, encoding="utf-8")
    browser.page.click("#next")
    browser.page.wait_for_timeout(500)
    (inbox / "house" / "page2.html").write_text(browser.page.content(), encoding="utf-8")

    # LGU: an Excel export with a title row above the header.
    (inbox / "lgu").mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["NCR LGU Directory"])
    ws.append(["LGU", "Mayor", "Vice Mayor", "Address"])
    ws.append(["Pateros", "Juan Dela Cruz", "Maria Santos", "Pateros"])
    ws.append(["Makati City", "Pedro Reyes", "Ana Garcia", "Makati"])
    wb.save(inbox / "lgu" / "lgu.xlsx")

    # Barangay: CSV in <Region>/<Province>/<City> folders, one row per official.
    folder = inbox / "barangay" / "Region VII" / "Bohol" / "Tagbilaran City"
    folder.mkdir(parents=True)
    with open(folder / "cogon.csv", "w", newline="", encoding="utf-8") as fh:
        w = _csv.writer(fh)
        w.writerow(["Barangay", "Position", "Name"])
        w.writerow(["Cogon", "Punong Barangay", "Jose Bautista"])
        w.writerow(["Cogon", "Barangay Kagawad", "Rosa Mendoza"])

    monkeypatch.setattr(sys, "argv", ["import_saved.py"])
    import_saved.main()

    data = isolated_output / "data"
    house_doc = json.loads((data / "house.json").read_text())
    assert house_doc["count"] == 100  # 50 per saved page
    assert house_doc["method"] == "saved-copy"
    assert house_doc["records"][0]["source"] == "House of Representatives (saved copy)"
    assert house_doc["records"][0]["details"]["copied_on"]

    lgu_doc = json.loads((data / "lgu.json").read_text())
    assert {(r["name"], r["level"], r["lgu"]) for r in lgu_doc["records"]} >= {
        ("Juan Dela Cruz", "mayor", "Pateros"), ("Ana Garcia", "vice_mayor", "Makati City")}

    bohol = json.loads((data / "barangay" / "region-vii--bohol.json").read_text())
    assert {(r["name"], r["level"], r["lgu"], r["barangay"]) for r in bohol["records"]} == {
        ("Jose Bautista", "punong_barangay", "Tagbilaran City", "Cogon"),
        ("Rosa Mendoza", "kagawad", "Tagbilaran City", "Cogon")}

    manifest = json.loads((data / "manifest.json").read_text())
    counts = {d["key"]: d["count"] for d in manifest["datasets"]}
    assert counts["house"] == 100 and counts["lgu"] == 4 and counts["barangay"] == 2


def test_senate_achievements_from_bio():
    """Mirrors the Senate API's biography HTML (bullets in <p> blocks)."""
    html = ('<p><span>Juan Dela Cruz (born 1970) is a Filipino politician.</span></p><p><br></p>'
            '<p><span>• Best Mayor in the Region (2008)</span></p>'
            '<p><span>• Outstanding Young Men, TOYM (2011)</span></p>'
            '<p>Co-author of the following laws:</p><p>• - Best Anchor</p>'
            '<ul><li>Author, Free Tuition Act</li></ul>'
            '<p><span>Served two decades in public service.</span></p>')
    assert senate.achievements(html) == [
        "Author, Free Tuition Act", "Best Mayor in the Region (2008)", "Outstanding Young Men, TOYM (2011)",
        "Best Anchor"]
    bio = senate.short_bio(html)
    assert bio.startswith("Juan Dela Cruz (born 1970)") and "Best Mayor" not in bio


def test_wikipedia_rowspans_and_mapping():
    """Layouts mirror the live Wikipedia lists (Sept 2026)."""
    import wikipedia

    house_html = """
    <table class="wikitable sortable"><tr><th>Constituency</th><th>Portrait</th><th>Representative</th>
      <th colspan="2">Party</th><th>Bloc</th><th>Born</th><th>Prior experience</th><th>Took office</th></tr>
    <tr><td>Bohol–1st</td><td><img src="//upload.example/a.jpg"></td>
        <td><a href="/wiki/Juan_Dela_Cruz">Juan Dela Cruz</a><sup>[1]</sup></td>
        <td style="background:red"></td><td>Lakas<sup>[a]</sup></td><td rowspan="2">Majority</td>
        <td><span style="display:none">( 1978-10-06 )</span> October 6, 1978 (age 47)</td>
        <td>House of Representatives<hr>Mayor of Tagbilaran</td><td>June 30, 2025</td></tr>
    <tr><td>Aklan at-large</td><td></td><td><a href="/wiki/Maria_Santos">Maria Santos</a></td>
        <td></td><td>NUP</td><td>1970</td><td></td><td>June 30, 2022</td></tr>
    <tr><td>Example Party</td><td></td><td><a href="/wiki/Pedro_Reyes">Pedro Reyes</a></td>
        <td></td><td>Example Party</td><td>Minority</td><td></td><td></td><td>June 30, 2025</td></tr>
    <tr><td>Leyte–2nd</td><td></td><td>Vacant</td><td></td><td></td><td></td><td></td><td></td><td></td></tr>
    </table>"""
    recs = wikipedia.officials("house", house_html, "123")
    by = {r["name"]: r for r in recs}
    assert set(by) == {"Juan Dela Cruz", "Maria Santos", "Pedro Reyes"}  # vacant seat skipped
    juan = by["Juan Dela Cruz"]
    assert (juan["position"], juan["district"], juan["province"], juan["party"]) == (
        "District Representative", "Bohol–1st", "Bohol", "Lakas")
    assert juan["photo"] == "https://upload.example/a.jpg"
    assert juan["profile_url"] == "https://en.wikipedia.org/wiki/Juan_Dela_Cruz"
    assert juan["details"]["born"] == "October 6, 1978"
    assert juan["details"]["prior_experience"] == "House of Representatives; Mayor of Tagbilaran"
    assert by["Maria Santos"]["province"] == "Aklan" and by["Maria Santos"]["details"]["bloc"] == "Majority"
    assert by["Pedro Reyes"]["position"] == "Party-list Representative" and by["Pedro Reyes"]["province"] == ""
    assert all(r["source"] == "Wikipedia (unofficial)" for r in recs)

    gov_html = """<table class="wikitable"><tr><th>Province</th><th>Portrait</th><th>Governor</th>
      <th colspan="2">Party</th><th>Term</th></tr>
      <tr><td>Abra<br>( <a href="/wiki/x">list</a> )</td><td></td><td><a href="/wiki/Takit_Bersamin">Takit Bersamin</a><sup>[3]</sup></td>
      <td></td><td>PFP</td><td>1</td></tr></table>"""
    (g,) = wikipedia.officials("governors", gov_html, "1")
    assert (g["name"], g["level"], g["province"], g["party"]) == ("Takit Bersamin", "governor", "Abra", "PFP")

    mayor_html = """<table class="wikitable"><tr><th>Independent city or municipality</th><th>Portrait</th>
      <th>Mayor</th><th colspan="2">Party</th><th>Age</th><th>Prior experience</th></tr>
      <tr><td>Baguio ( list )</td><td></td><td><a href="/wiki/B_M">Benjamin Magalong</a></td><td></td><td>NPC</td>
      <td><span>( 1970-06-04 )</span> June 4, 1970 (age 56)</td><td><div>Businessman</div><div>House of Representatives</div></td></tr></table>
      <table class="wikitable"><tr><th>Independent city or municipality</th><th>Portrait</th>
      <th>Vice mayor</th><th colspan="2">Party</th></tr>
      <tr><td>Baguio</td><td></td><td><a href="/wiki/F_O">Faustino Olowan</a></td><td></td><td>PFP</td></tr></table>"""
    m, v = wikipedia.officials("mayors", mayor_html, "1")
    assert m["details"]["born"] == "June 4, 1970"
    assert m["details"]["prior_experience"] == "Businessman; House of Representatives"
    assert (m["level"], m["lgu"], v["level"], v["lgu"]) == ("mayor", "Baguio", "vice_mayor", "Baguio")


def test_bills_attach_by_code_and_name(isolated_output):
    import bills

    data = isolated_output / "data"
    data.mkdir(parents=True, exist_ok=True)
    senators = [
        {"name": "Raffy T. Tulfo", "lis_code": ""}, {"name": "Erwin T. Tulfo", "lis_code": ""},
        {"name": "Win Gatchalian", "lis_code": "GSHER"},
        {"name": "Juan Miguel F. Zubiri", "lis_code": "ZJMIG, ZMIGU"},
    ]
    (data / "senate.json").write_text(json.dumps({"records": senators}))
    monkey = [
        {"number": "SBN-2", "bill_number": 2, "title": "B", "date": "", "status": "Pending", "url": "https://web.senate.gov.ph/x",
         "authors": ["TRAFF", "TERWI"], "authors_raw": "Tulfo, Raffy T., Tulfo, Erwin T."},
        {"number": "SBN-1", "bill_number": 1, "title": "A", "date": "", "status": "Pending", "url": "",
         "authors": ["TRAFF"], "authors_raw": "Tulfo, Raffy T."},
        {"number": "SBN-3", "bill_number": 3, "title": "C", "date": "", "status": "Pending", "url": "",
         "authors": ["TERWI"], "authors_raw": "Tulfo, Erwin T."},
        {"number": "SBN-4", "bill_number": 4, "title": "D", "date": "", "status": "Pending", "url": "",
         "authors": ["GSHER"], "authors_raw": "Gatchalian, Win"},
        {"number": "SBN-5", "bill_number": 5, "title": "E", "date": "", "status": "Pending", "url": "",
         "authors": ["ZJMIG"], "authors_raw": "Zubiri, Juan Miguel F."},
    ]
    import common

    doc = bills.attach(monkey, "2025-10-16")
    by = {r["name"]: r for r in doc["records"]}
    assert [b["number"] for b in by["Raffy T. Tulfo"]["bills"]] == ["SBN-2", "SBN-1"]  # initials don't cross-match
    assert [b["number"] for b in by["Erwin T. Tulfo"]["bills"]] == ["SBN-3", "SBN-2"]
    # Co-authorship is per senator: the first-listed author is the main author.
    assert by["Raffy T. Tulfo"]["bills"][0]["coauthored"] is False
    assert by["Erwin T. Tulfo"]["bills"][1]["coauthored"] is True
    assert by["Win Gatchalian"]["bills_count"] == 1
    assert by["Juan Miguel F. Zubiri"]["bills_count"] == 1  # several codes in one field
    assert doc["bills_source"]["source_type"] == "public" and doc["bills_source"]["as_of"] == "2025-10-16"


def test_bills_terms_and_service(isolated_output):
    import bills
    import common

    assert bills.senate_terms([20]) == 1
    assert bills.senate_terms([18, 19, 20]) == 2            # elected 2019, re-elected 2025
    assert bills.senate_terms([9, 10, 11, 12, 15, 16, 17, 18, 20]) == 5
    assert bills.senate_terms([11, 12, 14, 15, 16, 17, 19, 20]) == 4

    data = isolated_output / "data"
    data.mkdir(parents=True, exist_ok=True)
    (data / "senate.json").write_text(json.dumps({"records": [
        {"id": "senate-x-legarda", "name": "Loren Legarda", "lis_code": "LLORE"},
        {"id": "senate-x-new", "name": "New Senator", "lis_code": "NEWSE"}]}))
    mk = lambda c, n, authors, status="Pending in the Committee": {
        "congress": c, "number": f"SBN-{n}", "bill_number": n, "title": f"T{n}", "date": f"{2000 + c}-01-01",
        "status": status, "status_date": "", "url": "", "authors": authors, "authors_raw": ""}
    sample = [mk(20, 1, ["LLORE"]), mk(20, 2, ["NEWSE", "LLORE"]), mk(17, 5, ["LLORE"], "Approved by the President of the Philippines"),
              mk(14, 9, ["OTHER", "LLORE"], "Lapsed into Law")]
    people = {"LLORE": {"senate": [11, 12, 14, 15, 16, 17, 19, 20], "house": [18]}}
    congresses = {c: {"start": y, "end": y + 3, "ordinal": ""} for c, y in [(11, 1998), (14, 2007), (17, 2016), (20, 2025)]}
    doc = bills.attach(sample, "2025-10-16", people, congresses)
    by = {r["name"]: r for r in doc["records"]}
    leg = by["Loren Legarda"]
    assert leg["service"]["first_senate_year"] == 1998 and leg["service"]["senate_terms"] == 4
    assert leg["service"]["house_congresses"] == [18] and leg["service"]["basis"] == "membership"
    assert leg["bills_count"] == 2 and leg["bills_all"] == 4 and leg["law_all"] == 2
    t = {x["congress"]: x for x in leg["bill_terms"]}
    assert (t[20]["filed"], t[20]["main"]) == (2, 1) and t[17]["law"] == 1 and t[14]["main"] == 0
    assert 11 not in t  # no bill data before the 13th Congress
    new = by["New Senator"]["service"]
    assert new["basis"] == "bills" and new["senate_terms"] == 1 and new["first_senate_year"] == 2025
    page = json.loads((common.DATA_DIR / "bills" / "senate-x-legarda.json").read_text())
    assert [x["congress"] for x in page["terms"]] == [20, 19, 17, 16, 15, 14]
    assert page["terms"][2]["bills"][0]["law"] is True and page["terms"][0]["years"] == "2025–2028"

    # Sectors come from the committee referral; unknown committees are "other".
    assert bills.sector_of("Agriculture, Food and Agrarian Reform") == "agriculture"
    assert bills.sector_of("Basic Education, Arts and Culture") == "education"
    assert bills.sector_of("Cultural Communities and Muslim Affairs") == "social"
    assert bills.sector_of("Public Services") == "transport"
    assert bills.sector_of("") == "other"
    assert leg["bill_sectors"]["career"]["other"][0] == 4 and leg["bill_sectors"]["term"]["other"] == [2, 1]


def test_source_type():
    assert common.source_type("https://senate.gov.ph/senators") == "official"
    assert common.source_type("https://www.dilg.gov.ph/x") == "official"
    assert common.source_type("https://en.wikipedia.org/wiki/X") == "public"
    assert common.source_type("https://github.com/bettergovph/open-congress-data") == "public"
    assert common.source_type("https://notgov.ph.example.com/") == "public"


def test_dilg_directory_parse(isolated_output):
    """Layout guesses for the DILG sheet: title rows, split names, fill-down."""
    import io as _io

    import dilg_lgu
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "BOHOL"
    ws.append(["DIRECTORY OF LGU ELECTIVE OFFICIALS"])
    ws.append(["TERM 2025-2028"])
    ws.append(["REGION", "PROVINCE", "CITY/MUNICIPALITY", "POSITION", "LAST NAME", "FIRST NAME", "MIDDLE NAME", "SUFFIX"])
    ws.append(["VII", "BOHOL", "", "GOVERNOR", "AUMENTADO", "ARIS", "C", ""])
    ws.append(["", "", "TAGBILARAN CITY", "CITY MAYOR", "YAP", "JANE", "", ""])
    ws.append(["", "", "", "CITY VICE MAYOR", "CRUZ", "PEDRO", "D", "JR."])
    ws.append(["", "", "", "SANGGUNIANG PANLUNGSOD MEMBER", "REYES", "ANA", "", ""])
    ws.append(["", "", "PANGLAO", "MUNICIPAL MAYOR", "GARCIA", "JOSE", "", ""])
    ws.append(["", "", "", "SB MEMBERS", "LIM", "ROSA", "", ""])
    ws.append(["", "HUC", "CAGAYAN DE ORO CITY", "CITY MAYOR", "UY", "ROLANDO", "A", ""])
    ws.append(["", "", "", "SP MEMBERS (D1)", "DAHINO", "DESIREE", "", ""])
    ws.append(["", "", "", "", "ABADAY", "ROGER", "G", ""])
    ws.append(["", "BUKIDNON", "", "SP MEMBERS", "ALBARECE", "MARIO", "B", "JR."])
    buf = _io.BytesIO()
    wb.save(buf)
    recs = dilg_lgu.parse(buf.getvalue())
    got = {(r["name"], r["level"], r["province"], r["lgu"]) for r in recs}
    assert got == {
        ("Aris C. Aumentado", "governor", "Bohol", ""),
        ("Jane Yap", "mayor", "Bohol", "Tagbilaran City"),
        ("Pedro D. Cruz Jr.", "vice_mayor", "Bohol", "Tagbilaran City"),
        ("Ana Reyes", "councilor", "Bohol", "Tagbilaran City"),
        ("Jose Garcia", "mayor", "Bohol", "Panglao"),
        ("Rosa Lim", "councilor", "Bohol", "Panglao"),
        ("Rolando A. Uy", "mayor", "", "Cagayan de Oro City"),
        ("Desiree Dahino", "councilor", "", "Cagayan de Oro City"),
        ("Roger G. Abaday", "councilor", "", "Cagayan de Oro City"),
        ("Mario B. Albarece Jr.", "board_member", "Bukidnon", ""),
    }
    assert all(r["source_type"] == "official" for r in recs)


def test_dilg_picks_directory_not_folder(monkeypatch):
    import dilg_lgu

    listed = [(dilg_lgu.FOLDER_ID, "List of Elective Officials Term 2025-2028"),
              ("1B2N5SjZEBP-29tUioFHaSk-6R_1YFGcYIkSPjuApxZw", "DIRECTORY OF LGU ELECTIVE OFFICIALS FOR TERM 2025-2028")]
    picked = []
    monkeypatch.setattr(dilg_lgu, "find_files", lambda browser: [f for f in listed if f[0] != dilg_lgu.FOLDER_ID])
    monkeypatch.setattr(dilg_lgu, "download", lambda fid: picked.append(fid) or b"")
    monkeypatch.setattr(dilg_lgu, "parse", lambda data: [])
    dilg_lgu.run(browser=object())
    assert picked == ["1B2N5SjZEBP-29tUioFHaSk-6R_1YFGcYIkSPjuApxZw"]


def test_import_dilg_from_inbox(monkeypatch, isolated_output):
    import import_saved
    import run_all
    from openpyxl import Workbook

    inbox = isolated_output / "data-inbox"
    (inbox / "dilg").mkdir(parents=True)
    monkeypatch.setattr(import_saved, "INBOX", inbox)
    monkeypatch.setattr(run_all, "DATA_DIR", isolated_output / "data")
    import dilg_lgu
    monkeypatch.setattr(dilg_lgu, "MIN_OFFICIALS", 1)
    wb = Workbook()
    ws = wb.active
    ws.append(["PROVINCE", "CITY/MUNICIPALITY", "POSITION", "LAST NAME", "FIRST NAME"])
    ws.append(["BOHOL", "PANGLAO", "MUNICIPAL MAYOR", "GARCIA", "JOSE"])
    wb.save(inbox / "dilg" / "directory.xlsx")
    monkeypatch.setattr(sys, "argv", ["import_saved.py", "--only", "dilg"])
    import_saved.main()
    doc = json.loads((isolated_output / "data" / "lgu.json").read_text())
    (rec,) = doc["records"]
    assert (rec["name"], rec["level"], rec["lgu"], rec["source_type"]) == ("Jose Garcia", "mayor", "Panglao", "official")
    assert rec["details"]["copied_on"]


def test_dilg_region_import(site, browser, monkeypatch, isolated_output):
    import io as _io

    import dilg_lgu
    import dilg_regions
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "BUKIDNON"
    ws.append(["CITY/MUNICIPALITY", "POSITION", "NAME"])
    for i in range(25):
        ws.append([f"TOWN {i}", "MUNICIPAL MAYOR", f"JUAN DELA CRUZ {chr(65 + i)}"])
        ws.append([f"TOWN {i}", "MUNICIPAL VICE MAYOR", f"MARIA SANTOS {chr(65 + i)}"])
        for j in range(8):
            ws.append([f"TOWN {i}", "SANGGUNIANG BAYAN MEMBER", f"PEDRO REYES {chr(65 + i)}{chr(65 + j)}"])
    buf = _io.BytesIO()
    wb.save(buf)
    requested = []

    def fake_download(sid):
        requested.append(sid)
        if sid.startswith("1NotAllowed"):
            raise SystemExit("download not permitted")
        return buf.getvalue()

    monkeypatch.setattr(dilg_lgu, "download", fake_download)
    monkeypatch.setattr(dilg_regions, "REGION_PAGES", {"Region X (Northern Mindanao)": f"{site}/region.html"})
    counts = dilg_regions.run(browser)
    assert counts == {"Region X (Northern Mindanao)": 250}
    assert len(requested) == 2  # the refused sheet is skipped, not worked around
    doc = json.loads((isolated_output / "data" / "lgu.json").read_text())
    rec = next(r for r in doc["records"] if r["level"] == "mayor")
    assert rec["province"] == "Bukidnon" and rec["region"] == "Region X (Northern Mindanao)"
    assert rec["source"] == "DILG Region X (Northern Mindanao) – Local Officials 2025–2028"
    # Served from 127.0.0.1 here; real DILG pages are *.dilg.gov.ph -> "official".
    assert common.source_type("https://region10.dilg.gov.ph/local-officials/") == "official"
    assert "DILG Region X" in doc["source_label"]


def test_openhalalan_winners_terms_and_merge(isolated_output):
    import common
    import openhalalan

    cols = "Last Name,First Name,Middle Name,Middle Name Source,Title,Full Name,Position,Party,Year,Province,City,Region,Sex,Sex Source"
    rows = [
        # Three straight wins: term-limited in 2028.
        *[f'CRUZ,JUAN,SANTOS,x,,"CRUZ, JUAN SANTOS",MAYOR,LAKAS,{y},CAMARINES SUR,NAGA,REGION V,M,x' for y in (2019, 2022, 2025)],
        # Father then son with the same first name but a different middle name.
        'REYES,PEDRO,LOPEZ,x,,"REYES, PEDRO LOPEZ",MAYOR,NP,2022,CEBU,NAGA,REGION VII,M,x',
        'REYES,PEDRO,GARCIA,x,,"REYES, PEDRO GARCIA",MAYOR,NP,2025,CEBU,NAGA,REGION VII,M,x',
        'BELMONTE,JOY,G,x,,"BELMONTE, JOY G",MAYOR,SBP,2025,NCR SECOND DISTRICT,QUEZON,NATIONAL CAPITAL REGION,F,x',
        'SANTOS JR,ANA,,x,,"SANTOS JR, ANA",COUNCILOR,IND,2025,CEBU,NAGA,REGION VII,F,x',
        'UY,KLAREX,,x,,"UY, KLAREX",MAYOR,PFP,2025,MISAMIS ORIENTAL,CAGAYAN DE ORO,REGION X,M,x',
    ]
    found = openhalalan.officials(openhalalan.load(("\n".join([cols, *rows]) + "\n").encode()))
    by = {(o["name"], o["lgu"], o["province"]): o for o in found}
    assert by[("Juan S. Cruz", "Naga", "Camarines Sur")]["details"]["term"] == "3"
    assert by[("Pedro G. Reyes", "Naga", "Cebu")]["details"]["term"] == "1"
    assert ("Joy G. Belmonte", "Quezon City", "Metro Manila") in by
    ana = by[("Ana Santos Jr.", "Naga", "Cebu")]
    assert ana["party"] == "Independent" and ana["source_type"] == "public" and ana["level"] == "councilor"

    data = isolated_output / "data"
    data.mkdir(parents=True, exist_ok=True)
    existing = [
        # Official regional list: its whole region is left to it.
        {"level": "mayor", "name": "Rolando A. Uy", "province": "", "lgu": "Cagayan de Oro City",
         "region": "Region X (Northern Mindanao)", "source": "DILG Region X", "source_type": "official", "details": {}},
        # Wikipedia city mayor without a province: covers Naga, Camarines Sur only.
        {"level": "mayor", "name": "Leni Robredo", "province": "", "lgu": "Naga", "region": "",
         "source": "Wikipedia (unofficial)", "source_type": "public", "details": {}},
    ]
    (data / "lgu.json").write_text(json.dumps({"records": existing}))
    counts = openhalalan.merge(found)
    assert counts == {"added": 3, "in_official_regions": 1, "covered_by_other_sources": 1}
    doc = json.loads((common.DATA_DIR / "lgu.json").read_text())
    names = {r["name"] for r in doc["records"]}
    assert {"Pedro G. Reyes", "Joy G. Belmonte", "Leni Robredo", "Rolando A. Uy"} <= names
    assert "Juan S. Cruz" not in names and "Klarex Uy" not in names
    uy = next(r for r in doc["records"] if r["name"] == "Rolando A. Uy")
    assert uy["details"]["term"] == "1"  # term copied from the 2025 winner with the same surname


def test_dilg_barangay_sheet(isolated_output):
    """A regional punong barangay sheet becomes per-province barangay files."""
    import io as _io

    import barangay
    import common
    import dilg_barangays
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "BUKIDNON"
    ws.append(["LIST OF PUNONG BARANGAYS TERM 2023-2025"])
    ws.append(["PROVINCE", "CITY/MUNICIPALITY", "BARANGAY", "PUNONG BARANGAY", "CONTACT NO."])
    ws.append(["BUKIDNON", "BAUNGON", "BALINTAD", "JUAN DELA CRUZ", "0917"])
    ws.append(["", "", "IMBATUG", "MARIA SANTOS", ""])
    ws.append(["", "MALAYBALAY CITY", "CASISANG", "PEDRO REYES", ""])
    ws3 = wb.create_sheet("AURORA")  # DILG Central Office's long format
    ws3.append(["TERM", "REGION", "PROVINCE", "CITY/MUNICIPALITY", "BARANGAY", "POSITION", "TERM IN PRESENT POSITION",
                "LASTNAME", "FIRSTNAME", "MIDDLENAME", "SUFFIX"])
    ws3.append(["2023 - 2026", "REGION 3", "AURORA", "BALER", "Barangay I", "Punong Barangay", "3RD", "ZAFRA", "ALEJANDRO", "", "JR"])
    ws3.append(["2023 - 2026", "REGION 3", "AURORA", "BALER", "Barangay I", "Sangguniang Kabataan Member", "1ST", "CRUZ", "LEA", "", ""])
    ws3.append(["2023 - 2026", "REGION 3", "AURORA", "BALER", "Barangay I", "SK Treasurer", "", "DIAZ", "MAE", "", ""])
    ws2 = wb.create_sheet("Sheet1")
    ws2.append(["CITY/MUNICIPALITY", "BARANGAY", "PUNONG BARANGAY"])
    ws2.append(["CITY OF CAGAYAN DE ORO (Capital)", "CARMEN", "ana lim"])
    buf = _io.BytesIO()
    wb.save(buf)
    recs = dilg_barangays.parse(buf.getvalue(), "Region X (Northern Mindanao)", "https://region10.dilg.gov.ph/x", "SID")
    got = {(r["name"], r["level"], r["barangay"], r["lgu"], r["province"]) for r in recs}
    assert got == {
        ("Juan Dela Cruz", "punong_barangay", "Balintad", "Baungon", "Bukidnon"),
        ("Maria Santos", "punong_barangay", "Imbatug", "Baungon", "Bukidnon"),
        ("Pedro Reyes", "punong_barangay", "Casisang", "Malaybalay City", "Bukidnon"),
        ("Ana Lim", "punong_barangay", "Carmen", "Cagayan de Oro City", ""),
        ("Alejandro Zafra Jr", "punong_barangay", "Barangay I", "Baler", "Aurora"),
    }  # SK members and the SK treasurer are not council members
    by = {r["name"]: r for r in recs}
    assert by["Alejandro Zafra Jr"]["details"]["term"] == "3"
    # The national file: each row's region column is used.
    national = dilg_barangays.parse(buf.getvalue(), None, "https://www.dilg.gov.ph/x")
    assert {r["region"] for r in national if r["province"] == "Aurora"} == {"Region III (Central Luzon)"}
    assert all(r["source_type"] == "official" and r["dataset"] == "barangay" for r in recs)
    assert all(r["region"] == "Region X (Northern Mindanao)" and r["contact"] == "" for r in recs)
    barangay.flush(recs)
    index = json.loads((common.DATA_DIR / "barangay" / "index.json").read_text())
    assert index["count"] == 5 and {f["province"] for f in index["files"]} == {"Bukidnon", "Aurora", ""}


def test_executive_from_wikidata(isolated_output):
    import common
    import executive

    data = isolated_output / "data"
    data.mkdir(parents=True, exist_ok=True)
    (data / "laws.json").write_text(json.dumps({"as_of": "2025-10-16", "laws": [
        {"ra": "11900", "date": "2022-07-25", "lapsed": True},
        {"ra": "11934", "date": "2022-10-10", "lapsed": False},
        {"ra": "11800", "date": "2022-05-01", "lapsed": False},  # previous President
    ]}))
    holders = [
        {"officeLabel": "President of the Philippines", "person": "http://www.wikidata.org/entity/Q1",
         "personLabel": "Juan Cruz", "start": "2022-06-30T00:00:00Z"},
        {"officeLabel": "Vice President of the Philippines", "person": "http://www.wikidata.org/entity/Q2",
         "personLabel": "Maria Santos", "start": "2022-06-30T00:00:00Z"},
        {"officeLabel": "President of the Philippines", "person": "http://www.wikidata.org/entity/Q3",
         "personLabel": "Pedro Reyes", "start": "2016-06-30T00:00:00Z", "end": "2022-06-30T00:00:00Z", "ordinal": "16"},
        {"officeLabel": "President of the Philippines", "person": "http://www.wikidata.org/entity/Q4",
         "personLabel": "Old Leader", "start": "1961-12-30T00:00:00Z", "end": "1965-12-30T00:00:00Z",
         "died": "1990-01-01T00:00:00Z"},
    ]
    positions = {
        "Q1": [{"posLabel": "Senator of the Philippines", "start": "2010-06-30T00:00:00Z", "end": "2016-06-30T00:00:00Z"},
               {"posLabel": "President of the Philippines", "start": "2022-06-30T00:00:00Z"},
               {"posLabel": "Q123456"}],  # unlabelled item: skipped
        "Q2": [{"posLabel": "Mayor of Davao City", "start": "2016-06-30T00:00:00Z", "end": "2022-06-30T00:00:00Z"},
               {"posLabel": "Vice President of the Philippines", "start": "2022-06-30T00:00:00Z"}],
    }
    people = {"Q1": {"partyLabel": "PFP", "image": "http://commons.wikimedia.org/wiki/Special:FilePath/X.jpg"}, "Q2": {}}
    laws, as_of = executive.load_laws()
    summaries = {"Q1": "Juan Cruz is a Filipino politician."}
    built = executive.build(holders, positions, people, summaries, laws, as_of)
    assert [r["name"] for r in built] == ["Juan Cruz", "Old Leader", "Pedro Reyes", "Maria Santos"]
    past = {r["name"]: r for r in built if not r["current"]}
    assert past["Pedro Reyes"]["ordinal"] == "16" and past["Pedro Reyes"]["details"]["term"] == "2016–2022"
    assert past["Pedro Reyes"]["laws_signed"] == 1 and past["Pedro Reyes"]["laws_note"] == "partial"
    assert past["Old Leader"]["laws_note"] == "before_data" and "laws_signed" not in past["Old Leader"]
    recs = {r["level"]: r for r in built if r["current"]}
    assert recs["president"]["summary"].startswith("Juan Cruz") and "summary" not in recs["vice_president"]
    p, vp = recs["president"], recs["vice_president"]
    assert p["name"] == "Juan Cruz" and p["party"] == "PFP" and p["photo"].endswith("?width=300")
    assert p["details"]["term"] == "2022–2028" and p["source_type"] == "public"
    assert p["details"]["prior_experience"] == "Senator of the Philippines (2010–2016)"
    assert (p["laws_signed"], p["laws_lapsed"]) == (1, 1)
    assert (p["ra_first"], p["ra_last"], p["laws_est"], p["laws_coverage"]) == (11900, 11934, 35, 6)
    assert vp["successive_terms"] == 1 and "laws_signed" not in vp
    assert [i["title"] for i in p["career"]] == ["Senator of the Philippines", "President of the Philippines"]
    two = executive.career([{"posLabel": "Vice President of the Philippines", "start": "2016-06-30", "end": "2022-06-30"},
                            {"posLabel": "Vice President of the Philippines", "start": "2022-06-30"}])
    assert executive.successive_terms(two, "Vice President of the Philippines") == 2
