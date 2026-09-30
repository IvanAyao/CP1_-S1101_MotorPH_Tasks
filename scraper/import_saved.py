"""Build datasets from pages and files saved by hand from the official sites.

Some sources block automated access (Cloudflare) or only answer requests
from the Philippines. A person can still open them in a normal browser and
save them; this script turns those saved copies into the app's datasets.

Put files under data-inbox/ at the repo root:

    data-inbox/house/     saved pages of congress.gov.ph/house-members
    data-inbox/lgu/       saved pages / exports of the PSA LGU directory
    data-inbox/dilg/      DILG "Directory of LGU Elective Officials 2025-2028"
                          spreadsheet (.xlsx/.csv), e.g. obtained via FOI
    data-inbox/barangay/  saved pages / exports of the DILG directory
                          (optionally in <Region>/<Province>/ subfolders)
    data-inbox/senate/    saved pages of the Senate list (normally not needed)

Accepted formats: .html/.htm ("Save page as"), .mhtml/.mht (Chrome
"Download page" / "Save as single file"), .xlsx, .csv, .json.

Every record is labelled as a saved copy of its official source with the
date it was imported. Datasets are only replaced when the saved files yield
records, so an empty or wrong upload never wipes existing data.
"""

from __future__ import annotations

import argparse
import csv
import email
import io
import json
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup

import barangay
import house
import lgu
import senate
from common import ROOT, Capture, all_table_rows, dedupe, log, norm_key, write_dataset
from roster import people_from_capture

INBOX = ROOT / "data-inbox"
PAGE_EXT = {".html", ".htm", ".mhtml", ".mht"}
TABLE_EXT = {".xlsx", ".csv"}


def read_mhtml(raw: bytes) -> str:
    """The main HTML part of an MHTML archive."""
    msg = email.message_from_bytes(raw)
    for part in msg.walk():
        if part.get_content_type() == "text/html":
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            return payload.decode(charset, errors="replace")
    return ""


def read_page(path: Path, source_url: str) -> Capture:
    raw = path.read_bytes()
    html = read_mhtml(raw) if path.suffix.lower() in {".mhtml", ".mht"} else raw.decode("utf-8", errors="replace")
    return Capture(url=source_url, html=html, json_docs=[])


def read_table_file(path: Path) -> list[dict[str, str]]:
    """Rows of a CSV or the first sheet(s) of an XLSX, keyed by normalised header."""
    if path.suffix.lower() == ".csv":
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        return [{norm_key(k): (v or "").strip() for k, v in row.items() if k}
                for row in csv.DictReader(io.StringIO(text))]
    from openpyxl import load_workbook

    rows: list[dict[str, str]] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = None
        for values in it:
            cells = ["" if v is None else str(v).strip() for v in values]
            if header is None:
                # First row with at least two non-empty cells is the header.
                if sum(1 for c in cells if c) >= 2:
                    header = [norm_key(c) or f"col{i}" for i, c in enumerate(cells)]
                continue
            if any(cells):
                rows.append(dict(zip(header, cells)))
    return rows


def rows_from(path: Path, source_url: str) -> list[dict[str, str]]:
    ext = path.suffix.lower()
    if ext in TABLE_EXT:
        return read_table_file(path)
    if ext == ".json":
        from common import flatten, iter_object_lists
        doc = json.loads(path.read_text(encoding="utf-8"))
        return [flatten(x) for _, items in iter_object_lists(doc) for x in items]
    if ext in PAGE_EXT:
        soup = BeautifulSoup(read_page(path, source_url).html, "lxml")
        return [r for t in all_table_rows(soup) for r in t]
    return []


def files(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in PAGE_EXT | TABLE_EXT | {".json"})


def people_from_file(path: Path, source_url: str) -> list[dict]:
    if path.suffix.lower() in PAGE_EXT:
        method, people = people_from_capture(read_page(path, source_url))
        log(f"  {path.relative_to(INBOX)}: {len(people)} people via {method}")
        return people
    rows = rows_from(path, source_url)
    people = []
    from common import looks_like_name, pick, strip_honorific
    for row in rows:
        name = pick(row, r"^(name|full_name|member|representative|senator|official)") or next(
            (v for v in row.values() if looks_like_name(v)), "")
        if name and looks_like_name(name):
            people.append({"name": strip_honorific(name), "fields": row, "photo": "", "profile_url": ""})
    log(f"  {path.relative_to(INBOX)}: {len(people)} people from table")
    return people


def import_house(today: str) -> bool:
    paths = files(INBOX / "house")
    if not paths:
        return False
    people = [p for path in paths for p in people_from_file(path, house.URL)]
    recs = house.to_records(people, source="House of Representatives (saved copy)", copied_on=today)
    write_dataset("house", dedupe(recs), house.URL, min_count=1, max_count=400,
                  extra_meta={"method": "saved-copy", "copied_on": today})
    return True


def import_senate(today: str) -> bool:
    paths = files(INBOX / "senate")
    if not paths:
        return False
    recs = []
    for path in paths:
        method, people = people_from_capture(read_page(path, senate.URL)) if path.suffix.lower() in PAGE_EXT else ("table", people_from_file(path, senate.URL))
        for p in people:
            from common import record
            recs.append(record(name=p["name"], position="Senator", level="senate", district="Nationwide",
                               photo=p["photo"], profile_url=p["profile_url"],
                               source="Senate of the Philippines (saved copy)", source_url=senate.URL,
                               details={"copied_on": today}))
    write_dataset("senate", dedupe(recs), senate.URL, min_count=20, max_count=26,
                  extra_meta={"method": "saved-copy", "copied_on": today})
    return True


def import_lgu(today: str) -> bool:
    paths = files(INBOX / "lgu")
    if not paths:
        return False
    rows = [r for path in paths for r in rows_from(path, lgu.URL)]
    log(f"lgu: {len(rows)} rows from {len(paths)} file(s)")
    lgus, officials = lgu.from_rows(rows, source=f"{lgu.SOURCE} (saved copy)")
    for o in officials:
        o["details"]["copied_on"] = today
    write_dataset("lgu", officials, lgu.URL, min_count=1,
                  extra_meta={"lgus": lgus, "lgu_count": len(lgus), "method": "saved-copy", "copied_on": today})
    return True


def import_dilg(today: str) -> bool:
    """DILG's Directory of LGU Elective Officials (xlsx/csv obtained via FOI)."""
    import dilg_lgu

    paths = [p for p in files(INBOX / "dilg") if p.suffix.lower() in {".xlsx", ".csv"}]
    if not paths:
        return False
    officials = []
    for path in paths:
        if path.suffix.lower() == ".xlsx":
            found = dilg_lgu.parse(path.read_bytes())
        else:
            from officials import rows_to_officials
            found = rows_to_officials(read_table_file(path), source=dilg_lgu.SOURCE,
                                      source_url=dilg_lgu.FOI_URL, defaults={})
        log(f"  {path.relative_to(INBOX)}: {len(found)} officials")
        officials += found
    for o in officials:
        o["details"]["copied_on"] = today
    dilg_lgu.merge_into_lgu(dedupe(officials))
    return True


def import_barangay(today: str) -> bool:
    from officials import rows_to_officials

    base = INBOX / "barangay"
    paths = files(base)
    if not paths:
        return False
    recs = []
    for path in paths:
        # Optional <Region>/<Province>/ folders give context the table may lack.
        parts = path.relative_to(base).parts[:-1]
        defaults = {"region": parts[0] if len(parts) > 0 else "",
                    "province": parts[1] if len(parts) > 1 else "",
                    "lgu": parts[2] if len(parts) > 2 else ""}
        rows = rows_from(path, barangay.URL)
        found = rows_to_officials(rows, source=f"{barangay.SOURCE} (saved copy)",
                                  source_url=barangay.URL, defaults=defaults)
        for o in found:
            o["details"]["copied_on"] = today
        log(f"  {path.relative_to(INBOX)}: {len(found)} officials")
        recs += found
    if not recs:
        raise SystemExit("barangay: no officials found in data-inbox/barangay; keeping previous data")
    barangay.flush(dedupe(recs))
    return True


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--only", nargs="*", choices=["senate", "house", "lgu", "dilg", "barangay"])
    args = ap.parse_args()
    today = date.today().isoformat()
    wanted = args.only or ["senate", "house", "lgu", "dilg", "barangay"]
    importers = {"senate": import_senate, "house": import_house, "lgu": import_lgu,
                 "dilg": import_dilg, "barangay": import_barangay}
    failures, done = [], []
    for key in wanted:
        try:
            if importers[key](today):
                done.append(key)
        except (Exception, SystemExit) as err:  # noqa: BLE001 - one bad upload shouldn't block the rest
            log(f"{key} import FAILED: {err}")
            failures.append(key)
    import run_all
    run_all.build_manifest()
    log(f"imported: {', '.join(done) or 'nothing'}")
    if failures:
        raise SystemExit(f"failed: {', '.join(failures)}")


if __name__ == "__main__":
    main()
