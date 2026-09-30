"""Elected LGU officials 2025-2028 from DILG's published directory (official).

DILG released the "Directory of LGU Elective Officials for Term 2025-2028"
through the government FOI portal as a publicly shared Google Drive folder:

  FOI request:  https://www.foi.gov.ph/agencies/dilg/list-of-elected-government-official-for-the-term-2025-2028-by-province/
  Drive folder: https://drive.google.com/drive/folders/1fcyeVj_aNF3mdcnDq8rvL8Zj9r77oLZw

We open the public folder, find the spreadsheet in it, and download it with
Google's standard public export (the same file the folder's "Download"
button gives). No login or access controls are involved.

The sheet's layout isn't documented, so every run prints each tab's first
rows to the log, and the dataset is only written when a plausible number of
officials is parsed (see MIN_OFFICIALS).
"""

from __future__ import annotations

import argparse
import io
import json
import re
import urllib.request
from pathlib import Path

from common import DATA_DIR, USER_AGENT, Browser, clean, dedupe, log, norm_key, write_dataset
from officials import rows_to_officials

FOI_URL = "https://www.foi.gov.ph/agencies/dilg/list-of-elected-government-official-for-the-term-2025-2028-by-province/"
FOLDER_URL = "https://drive.google.com/drive/folders/1fcyeVj_aNF3mdcnDq8rvL8Zj9r77oLZw"
SOURCE = "DILG – Directory of LGU Elective Officials 2025–2028 (via FOI)"
# 82 provinces + ~1,640 cities/municipalities: governors, vice governors,
# board members, mayors, vice mayors and councilors run well over 10,000.
MIN_OFFICIALS = 5000
FILE_ID = re.compile(r"^[A-Za-z0-9_-]{25,}$")


def find_files(browser: Browser) -> list[tuple[str, str]]:
    """(file id, name) of the items listed in the public folder."""
    browser.open(FOLDER_URL)
    items = browser.page.evaluate("""() => [...document.querySelectorAll('[data-id]')].map(el => ({
        id: el.getAttribute('data-id'), text: (el.innerText || '').trim().slice(0, 200)
    }))""")
    seen, out = set(), []
    for it in items:
        fid, text = it["id"], clean(it["text"])
        if fid and FILE_ID.match(fid) and fid not in seen and text:
            seen.add(fid)
            out.append((fid, text.split("\n")[0]))
    log(f"dilg: folder lists {len(out)} item(s): {out}")
    return out


def download(file_id: str) -> bytes:
    """Public export of a shared Google Sheet (or plain download of a file)."""
    for url in (f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx",
                f"https://drive.google.com/uc?export=download&id={file_id}"):
        req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        try:
            with urllib.request.urlopen(req, timeout=180) as resp:
                data = resp.read()
            if data[:2] == b"PK":  # xlsx is a zip
                log(f"dilg: downloaded {len(data):,} bytes from {url.split('?')[0]}")
                return data
            log(f"dilg: {url} did not return a spreadsheet ({data[:60]!r})")
        except Exception as err:  # noqa: BLE001 - try the next form
            log(f"dilg: {url} failed: {err}")
    raise SystemExit("dilg: could not download the directory spreadsheet")


def sheet_rows(data: bytes) -> list[tuple[str, list[dict[str, str]]]]:
    """(tab name, rows) per worksheet; header = first row with 3+ filled cells.

    Provincial tabs often put the province in a title row above the header,
    so any single-cell row above it is remembered as that tab's context.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        header: list[str] | None = None
        title_lines: list[str] = []
        rows: list[dict[str, str]] = []
        preview = []
        for values in ws.iter_rows(values_only=True):
            cells = ["" if v is None else clean(v) for v in values]
            if len(preview) < 8:
                preview.append(" | ".join(c[:25] for c in cells if c)[:300])
            filled = [c for c in cells if c]
            if header is None:
                if len(filled) >= 3:
                    header = [norm_key(c) or f"col{i}" for i, c in enumerate(cells)]
                elif filled:
                    title_lines.append(" ".join(filled))
                continue
            if len(filled) >= 2:
                row = dict(zip(header, cells))
                row["__tab"] = ws.title
                row["__title"] = " / ".join(title_lines)
                rows.append(row)
        log(f"dilg: tab {ws.title!r}: {len(rows)} rows; header={header}")
        for line in preview:
            log(f"   | {line}")
        out.append((ws.title, rows))
    return out


def fill_down(rows: list[dict[str, str]], keys: tuple[str, ...]) -> None:
    """Directories often print a province/LGU once, then leave it blank below."""
    last: dict[str, str] = {}
    for row in rows:
        for k in list(row):
            if any(re.search(p, k) for p in keys):
                if row[k]:
                    last[k] = row[k]
                elif k in last:
                    row[k] = last[k]


ROMAN = re.compile(r"\b(Ii|Iii|Iv|Vi|Vii|Viii|Ix)\b")


def tidy_case(text: str) -> str:
    """'ARIS C. AUMENTADO III' -> 'Aris C. Aumentado III' (only if all caps)."""
    if not text or text != text.upper():
        return text
    out = text.title()
    out = ROMAN.sub(lambda m: m.group(1).upper(), out)
    return re.sub(r"\bMc(\w)", lambda m: "Mc" + m.group(1).upper(), out)


def parse(data: bytes) -> list[dict]:
    officials = []
    for tab, rows in sheet_rows(data):
        fill_down(rows, (r"region", r"province", r"city|municipal|lgu", r"district"))
        # A tab named after a province (and no province column) gives context.
        defaults = {}
        if rows and not any("province" in k for k in rows[0]):
            defaults["province"] = clean(re.sub(r"(?i)province of", "", tab))
        found = rows_to_officials(rows, source=SOURCE, source_url=FOI_URL, defaults=defaults)
        log(f"dilg: tab {tab!r} -> {len(found)} officials")
        officials += found
    for o in officials:
        for key in ("name", "province", "lgu", "region", "district"):
            o[key] = tidy_case(o[key])
        if o["level"] != "other":
            o["details"].pop("position_raw", None)  # keep the raw title only when unrecognised
    return dedupe(officials)


def run(browser: Browser | None = None) -> list[dict]:
    own = browser is None
    browser = browser or Browser()
    try:
        files = find_files(browser)
    finally:
        if own:
            browser.close()
    sheets = [f for f in files if re.search(r"director|elective|official", f[1], re.I)] or files
    if not sheets:
        raise SystemExit("dilg: no files found in the public folder")
    officials = parse(download(sheets[0][0]))
    counts: dict[str, int] = {}
    for o in officials:
        counts[o["level"]] = counts.get(o["level"], 0) + 1
    log(f"dilg: {len(officials)} officials by level: {counts}")
    return officials


def merge_into_lgu(officials: list[dict]) -> None:
    """Official DILG records replace Wikipedia ones for the same office."""
    path = DATA_DIR / "lgu.json"
    old = json.loads(path.read_text()) if path.exists() else {"records": []}
    covered = {(o["level"], (o["lgu"] or o["province"]).lower()) for o in officials}
    keep = [r for r in old.get("records", [])
            if "dilg" not in r.get("source", "").lower()
            and (r["level"], (r.get("lgu") or r.get("province") or "").lower()) not in covered]
    write_dataset("lgu", dedupe(officials + keep), FOI_URL, min_count=MIN_OFFICIALS,
                  extra_meta={"method": "official+wikipedia" if keep else "official",
                              "source_label": SOURCE, "published_via": FOLDER_URL,
                              "lgus": old.get("lgus", [])})


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--file", help="parse a downloaded .xlsx instead of fetching")
    args = ap.parse_args()
    officials = parse(Path(args.file).read_bytes()) if args.file else run()
    merge_into_lgu(officials)


if __name__ == "__main__":
    main()
